
#!/usr/bin/env python3
"""Chi-square analysis of eight RS response labels from the first sheet of four workbooks."""

import argparse
import os
from itertools import combinations

import numpy as np
import pandas as pd
from scipy.stats import chi2_contingency
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


FILE_TO_MODEL = {
    "g_scored_1.xlsx": "GPT-5.4",
    "g_scored_2.xlsx": "Claude 4.6",
    "g_scored_3.xlsx": "Qwen-3.6",
    "g_scored_4.xlsx": "Gemma-4",
}
MODELS = list(FILE_TO_MODEL.values())
N = 360
LABEL_AXIS = {
    "excessive_certainty": "RS-Reinforcing",
    "agency_outsourcing": "RS-Reinforcing",
    "retention_engagement": "RS-Reinforcing",
    "maladaptive_belief_validation": "RS-Reinforcing",
    "appropriate_limitation": "Boundary-Maintaining",
    "encouraging_autonomous_appraisal": "Boundary-Maintaining",
    "pattern_recognition": "Boundary-Maintaining",
    "adaptive_emotion_regulation": "Boundary-Maintaining",
}


def stars(p):
    return "***" if p < .001 else "**" if p < .01 else "*" if p < .05 else "ns"


def holm(pvals):
    pvals = np.asarray(pvals, float)
    m = len(pvals)
    order = np.argsort(pvals)
    adj = np.empty(m)
    running = 0.0
    for rank, idx in enumerate(order):
        running = max(running, (m - rank) * pvals[idx])
        adj[idx] = min(running, 1.0)
    return adj


def _presence_bool(value, filename, label, row_number):
    if pd.isna(value):
        raise ValueError(f"{filename}, Sheet1 row {row_number}, {label}_present: empty flag")
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, np.integer, float, np.floating)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        value = value.strip().lower()
        if value in ("true", "yes", "1"):
            return True
        if value in ("false", "no", "0"):
            return False
    raise ValueError(
        f"{filename}, Sheet1 row {row_number}, {label}_present: "
        f"expected True/False (or Yes/No or 1/0), got {value!r}"
    )


def load_counts(indir):
    counts = {label: {} for label in LABEL_AXIS}
    order = list(LABEL_AXIS)
    first_post_ids = None

    for fname, model in FILE_TO_MODEL.items():
        path = os.path.join(indir, fname)
        if not os.path.isfile(path):
            raise FileNotFoundError(f"Missing model file: {path}")

        df = pd.read_excel(path, sheet_name=0)
        required = ["post_id"] + [f"{label}_present" for label in order]
        missing = [col for col in required if col not in df.columns]
        if missing:
            raise ValueError(f"{fname}, first sheet: missing columns {missing}")
        if len(df) != N:
            raise ValueError(f"{fname}, first sheet: expected {N} posts, found {len(df)}")
        if df["post_id"].isna().any() or df["post_id"].duplicated().any():
            raise ValueError(f"{fname}, first sheet: empty or duplicated post_id")

        post_ids = set(df["post_id"].astype(str))
        if first_post_ids is None:
            first_post_ids = post_ids
        elif first_post_ids != post_ids:
            raise ValueError(f"{fname}: post_id set differs from the first workbook")

        for label in order:
            flags = [
                _presence_bool(v, fname, label, i)
                for i, v in enumerate(df[f"{label}_present"], start=2)
            ]
            counts[label][model] = {"present": sum(flags)}

        print(f"Read {fname}: {len(df)} posts, {len(order)} labels from first sheet")

    return counts, LABEL_AXIS, order


def run_omnibus(counts, axis, order):
    rows, praw = [], []
    for lab in order:
        present = [counts[lab][m]["present"] for m in MODELS]
        absent = [N - p for p in present]
        chi2, p, dof, _ = chi2_contingency([present, absent], correction=False)
        V = float(np.sqrt(chi2 / (N * len(MODELS))))
        row = {"Label": lab, "Axis": axis[lab]}
        for m, pr in zip(MODELS, present):
            row[f"{m} %"] = round(pr / N * 100, 1)
        row.update({"chi2": round(chi2, 2), "df": int(dof),
                    "p_raw": p, "Cramers_V": round(V, 3)})
        rows.append(row)
        praw.append(p)
    padj = holm(praw)
    for row, pa in zip(rows, padj):
        row["p_Holm"] = pa
        row["Sig"] = stars(pa)
    cols = (["Label", "Axis"] + [f"{m} %" for m in MODELS]
            + ["chi2", "df", "p_raw", "p_Holm", "Cramers_V", "Sig"])
    return pd.DataFrame(rows)[cols]


def run_pairwise(counts, axis, order):
    pairs = list(combinations(MODELS, 2))
    rows = []
    for lab in order:
        praw = []
        tmp = []
        for a, b in pairs:
            pa_, pb_ = counts[lab][a]["present"], counts[lab][b]["present"]
            chi2, p, dof, _ = chi2_contingency(
                [[pa_, N - pa_], [pb_, N - pb_]], correction=False)
            tmp.append((a, b, pa_, pb_, chi2, p))
            praw.append(p)
        padj = holm(praw)
        for (a, b, pa_, pb_, chi2, p), pa in zip(tmp, padj):
            rows.append({
                "Label": lab, "Axis": axis[lab],
                "Comparison": f"{a} vs {b}",
                "% A": round(pa_ / N * 100, 1),
                "% B": round(pb_ / N * 100, 1),
                "Diff_pp": round((pa_ - pb_) / N * 100, 1),
                "chi2": round(chi2, 2),
                "p_raw": p, "p_Holm": pa, "Sig": stars(pa),
            })
    return pd.DataFrame(rows)


def build_counts_table(counts, axis, order):
    rows = []
    for lab in order:
        row = {"Label": lab, "Axis": axis[lab]}
        for m in MODELS:
            c = counts[lab][m]
            row[f"{m} present_n"] = c["present"]
            row[f"{m} %"] = round(c["present"] / N * 100, 1)
        rows.append(row)
    return pd.DataFrame(rows)


def style_sheet(ws, n_header=1, pct_like=(), sci_like=()):
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    body_font = Font(name="Calibri", size=11)
    for col in ws.iter_cols():
        for cell in col:
            cell.font = body_font
    for cell in ws[n_header]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 40)
    ws.freeze_panes = ws.cell(row=n_header + 1, column=1)
    headers = {c.value: c.column_letter for c in ws[n_header]}
    for name, letter in headers.items():
        if name in sci_like:
            for cell in ws[letter][1:]:
                cell.number_format = "0.00E+00"
        elif name in pct_like:
            for cell in ws[letter][1:]:
                cell.number_format = "0.0"


def write_xlsx(omni, pair, cnts, out_path):
    tmp = out_path + ".tmp.xlsx"
    with pd.ExcelWriter(tmp, engine="openpyxl") as xw:
        omni.to_excel(xw, sheet_name="Omnibus_chi2", index=False)
        pair.to_excel(xw, sheet_name="Pairwise_posthoc", index=False)
        cnts.to_excel(xw, sheet_name="Present_counts", index=False)

    wb = load_workbook(tmp)
    style_sheet(wb["Omnibus_chi2"],
                pct_like=[f"{m} %" for m in MODELS] + ["Cramers_V"],
                sci_like=["p_raw", "p_Holm"])
    style_sheet(wb["Pairwise_posthoc"],
                pct_like=["% A", "% B", "Diff_pp"],
                sci_like=["p_raw", "p_Holm"])
    style_sheet(wb["Present_counts"],
                pct_like=[f"{m} %" for m in MODELS])
    wb.save(tmp)
    os.replace(tmp, out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--indir", default=".", help="folder with g_scored_1..4.xlsx")
    ap.add_argument("-o", "--out", default="rs_chi2_results.xlsx")
    args = ap.parse_args()

    counts, axis, order = load_counts(args.indir)
    omni = run_omnibus(counts, axis, order)
    pair = run_pairwise(counts, axis, order)
    cnts = build_counts_table(counts, axis, order)
    write_xlsx(omni, pair, cnts, args.out)

    pd.set_option("display.width", 200)
    print("OMNIBUS (chi-square homogeneity, df=3, Holm across 8 labels)")
    print(omni.to_string(index=False))
    print(f"\nPairwise rows: {len(pair)}  |  saved -> {args.out}")


if __name__ == "__main__":
    main()
