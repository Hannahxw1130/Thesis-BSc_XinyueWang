#!/usr/bin/env python3
"""Cohen's and Fleiss' kappa summarized in one Excel worksheet."""

import itertools
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from sklearn.metrics import cohen_kappa_score
from statsmodels.stats.inter_rater import aggregate_raters, fleiss_kappa


def normalize(series):
    return series.astype(str).str.strip().str.capitalize()


def display_name(column):
    name = column.removesuffix("_rs_judgment")
    if name == "gpt":
        return "GPT-5.4"
    if name.startswith("human") and name[5:].isdigit():
        return f"Human {name[5:]}"
    return name


def compute_cohen(df, col_a, col_b):
    a, b = normalize(df[col_a]), normalize(df[col_b])
    valid = a.isin(["Yes", "No"]) & b.isin(["Yes", "No"])
    if not valid.any():
        return None
    a, b = a[valid], b[valid]
    agreed = int((a == b).sum())
    return float(cohen_kappa_score(a, b)), len(a), agreed


def compute_fleiss(df, human_cols):
    labels = df[human_cols].apply(normalize)
    valid = labels.isin(["Yes", "No"]).all(axis=1)
    if not valid.any():
        return None
    encoded = labels.loc[valid].apply(lambda col: col.map({"Yes": 0, "No": 1})).astype(int)
    table, _ = aggregate_raters(encoded.to_numpy())
    agreed = int((encoded.nunique(axis=1) == 1).sum())
    return float(fleiss_kappa(table)), len(encoded), agreed


def main():
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("data.xlsx")
    gpt_col = sys.argv[2] if len(sys.argv) > 2 else "gpt_rs_judgment"
    output = source.with_name(f"{source.stem}_kappa_results.xlsx")

    df = pd.read_excel(source, sheet_name=0).iloc[:100].copy()
    human_cols = [
        col for col in df.columns
        if col.startswith("human") and col.endswith("_rs_judgment")
    ]
    if gpt_col not in df.columns:
        raise ValueError(f"Missing GPT column: {gpt_col}")
    if not human_cols:
        raise ValueError("No human*_rs_judgment columns found")

    results = []
    comparisons = [(gpt_col, col) for col in human_cols]
    comparisons += list(itertools.combinations(human_cols, 2))

    for col_a, col_b in comparisons:
        result = compute_cohen(df, col_a, col_b)
        if result is None:
            print(f"Skipped: {col_a} vs {col_b}; no valid judgments")
            continue
        kappa, total, agreed = result
        results.append({
            "Comparison": f"{display_name(col_a)} vs {display_name(col_b)}",
            "Test": "Cohen's Kappa",
            "Valid rows": total,
            "Kappa": round(kappa, 3),
            "Agreements": agreed,
            "Disagreements": total - agreed,
        })

    if len(human_cols) >= 2:
        result = compute_fleiss(df, human_cols)
        if result is not None:
            kappa, total, agreed = result
            results.append({
                "Comparison": "All human annotators",
                "Test": "Fleiss' Kappa",
                "Valid rows": total,
                "Kappa": round(kappa, 3),
                "Agreements": agreed,
                "Disagreements": total - agreed,
            })

    if not results:
        raise ValueError("No valid comparisons found")

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(results).to_excel(writer, sheet_name="Summary", index=False)
        ws = writer.sheets["Summary"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 24
        widths = {"A": 30, "B": 18, "C": 14, "D": 12, "E": 16, "F": 18}
        for letter, width in widths.items():
            ws.column_dimensions[letter].width = width
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="23395D")
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Calibri", size=11)
                if cell.column >= 3:
                    cell.alignment = Alignment(horizontal="center")
            row[3].number_format = "0.000"

    for item in results:
        print(f"{item['Comparison']}: {item['Test']} = {item['Kappa']:.3f}")
    print(f"Saved to: {output}")


if __name__ == "__main__":
    main()
